"""Inspect a Total Result (TR) workbook: sheets, header row, course blocks, student rows.

Week 1 data-study tool (S1). It only reads the file and prints what it finds; it does not
parse marks. Use it on every TR file to fill in docs/data-study/TR_FORMATS.md.

    python analytics/scripts/explore/inspect_tr.py data/<TR file>.xlsx
    python analytics/scripts/explore/inspect_tr.py data/<TR file>.xlsx --sheet ALL

Old .xls files: save them as .xlsx first (Excel or LibreOffice).
"""
import argparse
import re
import sys
from collections import Counter

from openpyxl import load_workbook

# 24ESKCS001: 2-digit year, institute code ESK, 2-letter branch code, 3-digit serial
ROLL_RE = re.compile(r"^\d{2}ESK[A-Z]{2}\d{3}$")
# words that mark a header row in any of the known TR layouts
HEADER_HINTS = ("roll", "rollno", "roll no", "roll_no")
# a course block starts with a course-code column: SHNAME1 / PSHNAME1 / DSHNAME1 / course_code_01 / Course Code 1
COURSE_RE = re.compile(r"^(P|D)?SHNAME\d+$|^course[ _]code[ _]?\d+$", re.IGNORECASE)
SCAN_ROWS = 10  # header is expected within the first rows


def clean(v):
    return str(v).strip() if v is not None else ""


def find_header_row(rows):
    """Index (0-based) of the first row that has a roll-number column, or None."""
    for i, row in enumerate(rows[:SCAN_ROWS]):
        if any(clean(v).lower() in HEADER_HINTS for v in row):
            return i
    return None


def inspect_sheet(ws):
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    print(f"\n=== Sheet '{ws.title}': {ws.max_row} rows x {ws.max_column} columns")
    h = find_header_row(rows)
    if h is None:
        print("  no header row with a roll-number column in the first rows -> not a TR sheet?")
        return
    header = [clean(v) for v in rows[h]]
    print(f"  header on row {h + 1}; {sum(1 for x in header if x)} named columns")

    roll_col = next(i for i, x in enumerate(header) if x.lower() in HEADER_HINTS)
    # some layouts (department per-branch sheets) split the header over two rows:
    # look for course-code columns in the header row and in the row below it
    courses = []
    for r in rows[h:h + 2]:
        courses += [(i + 1, clean(x)) for i, x in enumerate(r) if COURSE_RE.match(clean(x))]
    if len(rows) > h + 1 and any(COURSE_RE.match(clean(x)) for x in rows[h + 1]):
        print(f"  second header row: row {h + 2} (data starts on row {h + 3})")
    courses.sort()
    print(f"  roll number column: {roll_col + 1} ('{header[roll_col]}')")
    print(f"  course blocks found: {len(courses)}")
    for col, name in courses:
        print(f"    col {col:>3}  {name}")

    students, other = 0, Counter()
    for row in rows[h + 1:]:
        v = clean(row[roll_col]) if roll_col < len(row) else ""
        if ROLL_RE.match(v):
            students += 1
        elif v:
            other[v] += 1
    print(f"  student rows (roll no. like 24ESKCS001): {students}")
    if other:
        sample = ", ".join(f"'{k}' x{n}" for k, n in other.most_common(8))
        print(f"  other non-empty values in the roll column (summary rows?): {sample}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("file", help="TR workbook (.xlsx)")
    ap.add_argument("--sheet", help="inspect only this sheet")
    args = ap.parse_args()

    if args.file.lower().endswith(".xls"):
        sys.exit("Old .xls format: save the file as .xlsx first.")
    wb = load_workbook(args.file, read_only=True, data_only=True)
    print(f"File: {args.file}")
    print(f"Sheets: {wb.sheetnames}")
    sheets = [wb[args.sheet]] if args.sheet else wb.worksheets
    for ws in sheets:
        inspect_sheet(ws)


if __name__ == "__main__":
    main()
