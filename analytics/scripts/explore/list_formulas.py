"""List the formulas in a manual result-analysis workbook, sheet by sheet.

Week 1 data-study tool (S1). For every sheet it prints the size, the number of merged
ranges, how many cells are formulas, and one example of each distinct formula pattern
(cell references are replaced by '#', so '=COUNTIF(P4:P382,"F")' copied into 50 columns
counts as one pattern). It also flags broken formulas (#REF!). Use it to fill in
docs/data-study/MANUAL_WORKBOOK_MAP.md.

    python analytics/scripts/explore/list_formulas.py "data/I Sem  2024-25 TR Result Analysis.xlsx"
    python analytics/scripts/explore/list_formulas.py <workbook>.xlsx --sheet "Section wise"

Old .xls files: save them as .xlsx first (Excel or LibreOffice).
"""
import argparse
import re
import sys
import warnings
from collections import OrderedDict

from openpyxl import load_workbook

# openpyxl warns about page headers/footers it cannot parse; they do not matter here
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


def pattern(formula: str) -> str:
    """'=COUNTIF(P4:P382,"F")' -> '=COUNTIF(#:#,"F")' (cell references removed)."""
    return re.sub(r"\$?\b[A-Z]{1,3}\$?\d+\b", "#", formula)


def inspect_sheet(ws):
    patterns = OrderedDict()  # pattern -> (first cell, formula, count)
    total = broken = 0
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                total += 1
                if "#REF!" in v:
                    broken += 1
                key = pattern(v)
                if key in patterns:
                    first, f, n = patterns[key]
                    patterns[key] = (first, f, n + 1)
                else:
                    patterns[key] = (cell.coordinate, v, 1)

    print(f"\n=== '{ws.title}': {ws.max_row} rows x {ws.max_column} cols, "
          f"{len(ws.merged_cells.ranges)} merged ranges, {total} formulas")
    if not total:
        print("  no formulas: every value on this sheet is typed")
        return
    if broken:
        print(f"  !! {broken} formulas contain #REF! (broken reference)")
    print(f"  {len(patterns)} distinct patterns (example cell, copies, formula):")
    for first, f, n in patterns.values():
        print(f"    {first:>7}  x{n:<4} {f[:110]}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("file", help="workbook (.xlsx)")
    ap.add_argument("--sheet", help="only this sheet")
    args = ap.parse_args()

    if args.file.lower().endswith(".xls"):
        sys.exit("Old .xls format: save the file as .xlsx first.")
    # data_only=False (default) keeps the formulas instead of their cached values
    wb = load_workbook(args.file)
    print(f"File: {args.file}")
    print(f"Sheets: {wb.sheetnames}")
    sheets = [wb[args.sheet]] if args.sheet else wb.worksheets
    for ws in sheets:
        inspect_sheet(ws)


if __name__ == "__main__":
    main()
